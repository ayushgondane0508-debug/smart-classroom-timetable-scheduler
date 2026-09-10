import asyncio
import logging
import os
import uuid
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import bcrypt
import jwt
from dotenv import load_dotenv
from fastapi import APIRouter, BackgroundTasks, Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, ConfigDict, EmailStr, Field

from emailer import email_configured, schedule_html, send_email
from scheduler import generate_schedule, validate_entries
from storage import get_object, init_storage, put_object

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")
client = AsyncIOMotorClient(os.environ["MONGO_URL"])
db = client[os.environ["DB_NAME"]]
app = FastAPI(title="Smart Classroom Scheduler API")
api = APIRouter(prefix="/api")
JWT_ALGORITHM = "HS256"


class LoginInput(BaseModel):
    email: EmailStr
    password: str = Field(min_length=6, max_length=128)


class ContactMessageCreate(BaseModel):
    name: str = Field(min_length=2, max_length=100)
    email: EmailStr
    message: str = Field(min_length=10, max_length=2000)


class EntityInput(BaseModel):
    model_config = ConfigDict(extra="allow")


class CredentialsInput(BaseModel):
    email: EmailStr
    password: str = Field(min_length=6, max_length=128)


class ChangePasswordInput(BaseModel):
    current_password: str = Field(min_length=6, max_length=128)
    new_password: str = Field(min_length=8, max_length=128)


MAX_PHOTO_BYTES = 3 * 1024 * 1024
MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
LOCKOUT_ATTEMPTS = 5
LOCKOUT_MINUTES = 15


def now():
    return datetime.now(timezone.utc).isoformat()


def hash_password(password):
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()


def verify_password(password, hashed):
    return bcrypt.checkpw(password.encode(), hashed.encode())


def make_token(user_id, email, role):
    return jwt.encode({"sub": user_id, "email": email, "role": role, "exp": datetime.now(timezone.utc) + timedelta(hours=8), "type": "access"}, os.environ["JWT_SECRET"], algorithm=JWT_ALGORITHM)


async def current_user(request: Request):
    token = request.cookies.get("access_token") or request.headers.get("Authorization", "").removeprefix("Bearer ")
    if not token:
        raise HTTPException(status_code=401, detail="Login required")
    try:
        payload = jwt.decode(token, os.environ["JWT_SECRET"], algorithms=[JWT_ALGORITHM])
    except jwt.InvalidTokenError as exc:
        raise HTTPException(status_code=401, detail="Invalid session") from exc
    user = await db.users.find_one({"id": payload.get("sub")}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(status_code=401, detail="Session expired")
    return user


async def admin_user(user: dict = Depends(current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    return user


async def teacher_user(user: dict = Depends(current_user)):
    if user.get("role") != "teacher":
        raise HTTPException(status_code=403, detail="Teacher account required")
    return user


async def check_lockout(identifier):
    record = await db.login_attempts.find_one({"identifier": identifier})
    if record and record.get("count", 0) >= LOCKOUT_ATTEMPTS and record.get("locked_until", "") > now():
        raise HTTPException(status_code=429, detail="Too many failed attempts. Try again in 15 minutes.")


async def record_failure(identifier):
    await db.login_attempts.update_one({"identifier": identifier}, {"$inc": {"count": 1}, "$set": {"locked_until": (datetime.now(timezone.utc) + timedelta(minutes=LOCKOUT_MINUTES)).isoformat()}}, upsert=True)


async def seed_admin():
    email = os.environ["ADMIN_EMAIL"].lower()
    existing = await db.users.find_one({"email": email})
    if not existing:
        await db.users.insert_one({"id": str(uuid.uuid4()), "email": email, "name": "Administrator", "role": "admin", "password_hash": hash_password(os.environ["ADMIN_PASSWORD"]), "created_at": now()})
    elif not verify_password(os.environ["ADMIN_PASSWORD"], existing["password_hash"]):
        await db.users.update_one({"email": email}, {"$set": {"password_hash": hash_password(os.environ["ADMIN_PASSWORD"])}})


@app.on_event("startup")
async def startup():
    await seed_admin()
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier")
    for collection in ["teachers", "subjects", "divisions", "classrooms", "laboratories", "files", "notifications"]:
        await db[collection].create_index("id", unique=True)
    try:
        await asyncio.to_thread(init_storage)
    except Exception as exc:
        logging.getLogger("storage").error("Storage init failed: %s", exc)


@api.get("/")
async def root():
    return {"message": "Smart Classroom Scheduler API"}


def public_user(user):
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"], "teacher_id": user.get("teacher_id")}


@api.post("/auth/login")
async def login(input: LoginInput, request: Request, response: Response):
    identifier = f"{request.client.host if request.client else 'unknown'}:{input.email.lower()}"
    await check_lockout(identifier)
    user = await db.users.find_one({"email": input.email.lower()})
    if not user or not verify_password(input.password, user["password_hash"]):
        await record_failure(identifier)
        raise HTTPException(status_code=401, detail="Incorrect email or password")
    await db.login_attempts.delete_one({"identifier": identifier})
    response.set_cookie("access_token", make_token(user["id"], user["email"], user["role"]), httponly=True, secure=True, samesite="none", max_age=28800, path="/")
    return public_user(user)


@api.post("/auth/logout")
async def logout(response: Response, _: dict = Depends(current_user)):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api.get("/auth/me")
async def me(user: dict = Depends(current_user)):
    return public_user(user)


@api.post("/auth/change-password")
async def change_password(input: ChangePasswordInput, user: dict = Depends(current_user)):
    record = await db.users.find_one({"id": user["id"]})
    if not verify_password(input.current_password, record["password_hash"]):
        raise HTTPException(status_code=400, detail="Current password is incorrect")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password_hash": hash_password(input.new_password), "updated_at": now()}})
    return {"ok": True}


@api.put("/teachers/{teacher_id}/credentials")
async def set_teacher_credentials(teacher_id: str, input: CredentialsInput, _: dict = Depends(admin_user)):
    teacher = await db.teachers.find_one({"id": teacher_id}, {"_id": 0})
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found")
    email = input.email.lower()
    clash = await db.users.find_one({"email": email, "teacher_id": {"$ne": teacher_id}})
    if clash:
        raise HTTPException(status_code=409, detail="That email is already used by another account")
    existing = await db.users.find_one({"teacher_id": teacher_id})
    if existing:
        await db.users.update_one({"teacher_id": teacher_id}, {"$set": {"email": email, "name": teacher["name"], "password_hash": hash_password(input.password), "updated_at": now()}})
    else:
        await db.users.insert_one({"id": str(uuid.uuid4()), "email": email, "name": teacher["name"], "role": "teacher", "teacher_id": teacher_id, "password_hash": hash_password(input.password), "created_at": now()})
    await db.teachers.update_one({"id": teacher_id}, {"$set": {"email": email, "has_login": True, "updated_at": now()}})
    return {"ok": True, "email": email, "login_url": f"{os.environ['FRONTEND_URL']}/teacher-login"}


@api.delete("/teachers/{teacher_id}/credentials")
async def revoke_teacher_credentials(teacher_id: str, _: dict = Depends(admin_user)):
    await db.users.delete_many({"teacher_id": teacher_id, "role": "teacher"})
    await db.teachers.update_one({"id": teacher_id}, {"$set": {"has_login": False, "updated_at": now()}})
    return {"ok": True}


@api.post("/teachers/{teacher_id}/photo")
async def upload_teacher_photo(teacher_id: str, file: UploadFile = File(...), _: dict = Depends(admin_user)):
    teacher = await db.teachers.find_one({"id": teacher_id}, {"_id": 0})
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found")
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(status_code=400, detail="Only image files are allowed")
    data = await file.read()
    if len(data) > MAX_PHOTO_BYTES:
        raise HTTPException(status_code=400, detail="Photo must be smaller than 3 MB")
    record = await store_file(data, file.filename, file.content_type, kind="photo", folder=f"photos/{teacher_id}", public=True, title=f"{teacher['name']} photo")
    await db.teachers.update_one({"id": teacher_id}, {"$set": {"photo_file_id": record["id"], "updated_at": now()}})
    return {"ok": True, "photo_file_id": record["id"]}


async def store_file(data, filename, content_type, kind, folder, public=False, title="", category="", meta=None):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "bin"
    file_id = str(uuid.uuid4())
    result = await put_object(f"{folder}/{file_id}.{ext}", data, content_type or "application/octet-stream")
    record = {"id": file_id, "kind": kind, "public": public, "title": title or filename, "category": category, "original_filename": filename, "content_type": content_type or "application/octet-stream", "size": result.get("size", len(data)), "storage_path": result["path"], "is_deleted": False, "created_at": now(), **(meta or {})}
    await db.files.insert_one(record)
    record.pop("_id", None)
    return record


@api.get("/files/{file_id}")
async def download_file(file_id: str, request: Request):
    record = await db.files.find_one({"id": file_id, "is_deleted": False}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=404, detail="File not found")
    if not record.get("public"):
        await current_user(request)
    try:
        data, content_type = await get_object(record["storage_path"])
    except Exception as exc:
        raise HTTPException(status_code=502, detail="Storage is temporarily unavailable") from exc
    disposition = "inline" if record.get("public") or record["content_type"].startswith("image/") else "attachment"
    return Response(content=data, media_type=record.get("content_type") or content_type, headers={"Content-Disposition": f'{disposition}; filename="{record["original_filename"]}"'})


@api.get("/documents")
async def list_documents(_: dict = Depends(current_user)):
    return await db.files.find({"kind": "document", "is_deleted": False}, {"_id": 0}).sort("created_at", -1).to_list(500)


@api.post("/documents")
async def upload_document(file: UploadFile = File(...), title: str = Form(""), category: str = Form("General"), _: dict = Depends(admin_user)):
    data = await file.read()
    if len(data) > MAX_DOCUMENT_BYTES:
        raise HTTPException(status_code=400, detail="Documents must be smaller than 15 MB")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty")
    try:
        return await store_file(data, file.filename, file.content_type, kind="document", folder="documents", title=title.strip(), category=category.strip() or "General")
    except Exception as exc:
        raise HTTPException(status_code=502, detail="Could not reach cloud storage. Please retry.") from exc


@api.delete("/documents/{file_id}")
async def delete_document(file_id: str, _: dict = Depends(admin_user)):
    result = await db.files.update_one({"id": file_id, "kind": "document"}, {"$set": {"is_deleted": True, "deleted_at": now()}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Document not found")
    return {"ok": True}


@api.get("/exports")
async def list_exports(_: dict = Depends(admin_user)):
    return await db.files.find({"kind": "export", "is_deleted": False}, {"_id": 0}).sort("created_at", -1).to_list(200)


async def archive_export(timetable, data, ext, content_type):
    try:
        await store_file(data, f"{timetable.get('name', 'timetable')}.{ext}", content_type, kind="export", folder=f"exports/{timetable['id']}", title=f"{timetable.get('name', 'Timetable')} ({ext.upper()})", meta={"timetable_id": timetable["id"], "timetable_name": timetable.get("name", ""), "score": timetable.get("score", 0)})
    except Exception as exc:
        logging.getLogger("storage").error("Export archive failed: %s", exc)


async def notify_teachers(timetable_id):
    timetable = await db.timetables.find_one({"id": timetable_id}, {"_id": 0})
    if not timetable:
        return []
    results = []
    for teacher in await list_entities("teachers"):
        entries = [entry for entry in timetable.get("entries", []) if entry.get("teacher_id") == teacher["id"]]
        if not entries or not teacher.get("email"):
            continue
        subject = f"New timetable published: {timetable.get('name', 'Weekly timetable')}"
        outcome = await send_email(teacher["email"], subject, schedule_html(teacher, timetable, entries, f"{os.environ['FRONTEND_URL']}/teacher/{teacher['id']}"))
        record = {"id": str(uuid.uuid4()), "timetable_id": timetable_id, "timetable_name": timetable.get("name", ""), "teacher_id": teacher["id"], "teacher_name": teacher["name"], "email": teacher["email"], "subject": subject, "sessions": len(entries), **outcome, "created_at": now()}
        await db.notifications.insert_one(record)
        record.pop("_id", None)
        results.append(record)
    return results


@api.get("/notifications")
async def list_notifications(_: dict = Depends(admin_user)):
    items = await db.notifications.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)
    return {"configured": email_configured(), "sender": os.environ.get("SENDER_EMAIL", ""), "items": items}


@api.post("/timetable/{timetable_id}/notify")
async def notify_now(timetable_id: str, _: dict = Depends(admin_user)):
    results = await notify_teachers(timetable_id)
    return {"ok": True, "configured": email_configured(), "count": len(results), "items": results}


@api.post("/contact")
async def contact(input: ContactMessageCreate):
    item = {"id": str(uuid.uuid4()), **input.model_dump(), "created_at": now()}
    await db.contact_messages.insert_one(item)
    item.pop("_id", None)
    return item


async def list_entities(collection):
    return await db[collection].find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)


def entity_router(name):
    @api.get(f"/{name}")
    async def get_all(_: dict = Depends(admin_user)):
        return await list_entities(name)

    @api.post(f"/{name}")
    async def create(input: EntityInput, _: dict = Depends(admin_user)):
        item = {"id": str(uuid.uuid4()), **input.model_dump(), "created_at": now(), "updated_at": now()}
        await db[name].insert_one(item)
        item.pop("_id", None)
        return item

    @api.put(f"/{name}/{{item_id}}")
    async def update(item_id: str, input: EntityInput, _: dict = Depends(admin_user)):
        item = {**input.model_dump(), "updated_at": now()}
        result = await db[name].update_one({"id": item_id}, {"$set": item})
        if result.matched_count == 0:
            raise HTTPException(status_code=404, detail="Record not found")
        return {"id": item_id, **item}

    @api.delete(f"/{name}/{{item_id}}")
    async def delete(item_id: str, _: dict = Depends(admin_user)):
        result = await db[name].delete_one({"id": item_id})
        if result.deleted_count == 0:
            raise HTTPException(status_code=404, detail="Record not found")
        return {"ok": True}


for entity in ["teachers", "subjects", "divisions", "classrooms", "laboratories"]:
    entity_router(entity)


DEFAULT_CONFIG = {"id": "default", "college_name": "Smart Classroom College", "academic_year": "2026", "department_mode": "single", "working_days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"], "periods_per_day": 6, "period_duration": 55, "start_time": "09:00", "break_after_period": 3, "lunch_after_period": 5}


@api.get("/config")
async def get_config(_: dict = Depends(admin_user)):
    return await read_config()


async def read_config():
    return await db.settings.find_one({"id": "default"}, {"_id": 0}) or DEFAULT_CONFIG


@api.put("/config")
async def update_config(input: EntityInput, _: dict = Depends(admin_user)):
    item = {"id": "default", **input.model_dump(), "updated_at": now()}
    await db.settings.replace_one({"id": "default"}, item, upsert=True)
    return item


def demo_documents():
    teachers = [{"id": f"demo-t-{i}", "name": name, "employee_id": f"FAC-{i:03}", "department": "Computer Engineering", "maximum_lectures_per_day": 4, "maximum_lectures_per_week": 20, "availability": [], "demo": True, "created_at": now()} for i, name in enumerate(["Aarav Shah", "Meera Joshi", "Kabir Patil", "Nisha Rao", "Rohan Kulkarni", "Isha Deshmukh"], 1)]
    teachers += [{"id": f"demo-t-{i}", "name": name, "employee_id": f"FAC-{i:03}", "department": "Information Technology", "departments": ["Computer Engineering"], "maximum_lectures_per_day": 4, "maximum_lectures_per_week": 20, "availability": [], "demo": True, "created_at": now()} for i, name in enumerate(["Priya Nair", "Dev Mehta", "Sana Khan"], 7)]
    subjects = [{"id": f"demo-s-{i}", "name": name, "code": code, "department": "Computer Engineering", "semester": 3, "type": kind, "lectures_per_week": count, "duration": 1, "requires_lab": lab, "teacher_id": f"demo-t-{((i - 1) % 6) + 1}", "demo": True, "created_at": now()} for i, (name, code, kind, count, lab) in enumerate([("Data Structures", "CS201", "Theory", 3, False), ("Operating Systems", "CS202", "Theory", 3, False), ("Database Systems", "CS203", "Theory", 3, False), ("Computer Networks", "CS204", "Theory", 2, False), ("DS Lab", "CS205L", "Practical", 2, True), ("DBMS Lab", "CS206L", "Practical", 2, True)], 1)]
    subjects += [{"id": f"demo-s-{i}", "name": name, "code": code, "department": "Information Technology", "semester": 3, "type": kind, "lectures_per_week": count, "duration": 1, "requires_lab": lab, "teacher_id": teacher, "demo": True, "created_at": now()} for i, (name, code, kind, count, lab, teacher) in enumerate([("Web Technologies", "IT301", "Theory", 3, False, "demo-t-7"), ("Software Engineering", "IT302", "Theory", 3, False, "demo-t-8"), ("Object Oriented Programming", "IT303", "Theory", 2, False, "demo-t-9"), ("Web Lab", "IT304L", "Practical", 2, True, "demo-t-7"), ("OOP Lab", "IT305L", "Practical", 2, True, "demo-t-9")], 7)]
    cse_subjects = [s["id"] for s in subjects if s["department"] == "Computer Engineering"]
    it_subjects = [s["id"] for s in subjects if s["department"] == "Information Technology"]
    divisions = [{"id": "demo-d-1", "name": "CSE-A", "department": "Computer Engineering", "semester": 3, "student_count": 42, "subjects": cse_subjects, "demo": True, "created_at": now()}, {"id": "demo-d-2", "name": "CSE-B", "department": "Computer Engineering", "semester": 3, "student_count": 36, "subjects": cse_subjects, "demo": True, "created_at": now()}, {"id": "demo-d-3", "name": "IT-A", "department": "Information Technology", "semester": 3, "student_count": 40, "subjects": it_subjects, "demo": True, "created_at": now()}, {"id": "demo-d-4", "name": "IT-B", "department": "Information Technology", "semester": 3, "student_count": 38, "subjects": it_subjects, "demo": True, "created_at": now()}]
    classrooms = [{"id": f"demo-r-{i}", "room_number": f"20{i}", "building": "A Block", "capacity": 60, "room_type": "Classroom", "available_slots": [], "demo": True, "created_at": now()} for i in range(1, 5)]
    labs = [{"id": "demo-l-1", "lab_name": "Systems Lab", "lab_number": "LAB-1", "lab_type": "Computer", "capacity": 50, "room_type": "Laboratory", "available_slots": [], "demo": True, "created_at": now()}, {"id": "demo-l-2", "lab_name": "Database Lab", "lab_number": "LAB-2", "lab_type": "Computer", "capacity": 50, "room_type": "Laboratory", "available_slots": [], "demo": True, "created_at": now()}, {"id": "demo-l-3", "lab_name": "Web Lab", "lab_number": "LAB-3", "lab_type": "Computer", "capacity": 50, "room_type": "Laboratory", "available_slots": [], "demo": True, "created_at": now()}]
    return {"teachers": teachers, "subjects": subjects, "divisions": divisions, "classrooms": classrooms, "laboratories": labs}


@api.post("/demo/load")
async def load_demo(_: dict = Depends(admin_user)):
    docs = demo_documents()
    for collection, values in docs.items():
        await db[collection].delete_many({"demo": True})
        await db[collection].insert_many(values)
    await db.settings.replace_one({"id": "default"}, DEFAULT_CONFIG, upsert=True)
    return {"ok": True, "counts": {name: len(values) for name, values in docs.items()}}


@api.delete("/demo/reset")
async def reset_demo(_: dict = Depends(admin_user)):
    for collection in ["teachers", "subjects", "divisions", "classrooms", "laboratories", "timetables"]:
        await db[collection].delete_many({"demo": True})
    return {"ok": True}


@api.get("/dashboard/stats")
async def dashboard_stats(_: dict = Depends(admin_user)):
    config = await read_config()
    counts = {name: await db[name].count_documents({}) for name in ["teachers", "subjects", "divisions", "classrooms", "laboratories"]}
    timetable_count = await db.timetables.count_documents({})
    current = await db.timetables.find_one({"active": True}, {"_id": 0})
    return {**counts, "weekly_periods": len(config.get("working_days", [])) * int(config.get("periods_per_day", 0)), "generated_timetables": timetable_count, "current_conflicts": len(validate_entries(current.get("entries", []))) if current else 0, "faculty_workload": [], "room_utilization": [], "subject_distribution": []}


@api.post("/timetable/generate")
async def generate(background: BackgroundTasks, _: dict = Depends(admin_user)):
    data = {name: await list_entities(name) for name in ["teachers", "subjects", "divisions", "classrooms", "laboratories"]}
    data["config"] = await read_config()
    result = generate_schedule(teachers=data["teachers"], subjects=data["subjects"], divisions=data["divisions"], rooms=data["classrooms"], labs=data["laboratories"], config=data["config"])
    if result["failures"]:
        return {"ok": False, "message": "Unable to generate a completely conflict-free timetable.", "failures": result["failures"], "score": result["score"], "hard_constraint_violations": result["hard_constraint_violations"]}
    timetable = {"id": str(uuid.uuid4()), "name": f"{data['config'].get('college_name', 'College')} — {data['config'].get('academic_year', 'Current')}", "entries": result["entries"], "slots": result["slots"], "score": result["score"], "hard_constraint_violations": 0, "active": True, "created_at": now(), "updated_at": now()}
    await db.timetables.update_many({}, {"$set": {"active": False}})
    await db.timetables.insert_one(timetable)
    timetable.pop("_id", None)
    background.add_task(notify_teachers, timetable["id"])
    return {"ok": True, "timetable": timetable}


@api.get("/timetable")
async def get_timetable(_: dict = Depends(admin_user)):
    return await db.timetables.find({}, {"_id": 0}).sort("created_at", -1).to_list(100)


@api.put("/timetable/{timetable_id}")
async def update_timetable(timetable_id: str, input: EntityInput, _: dict = Depends(admin_user)):
    payload = input.model_dump(); conflicts = validate_entries(payload.get("entries", []))
    if conflicts:
        raise HTTPException(status_code=409, detail={"message": "Conflict detected", "conflicts": conflicts})
    payload["updated_at"] = now()
    result = await db.timetables.update_one({"id": timetable_id}, {"$set": payload})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Timetable not found")
    return {"id": timetable_id, **payload}


@api.get("/analytics")
async def analytics(_: dict = Depends(admin_user)):
    return await compute_analytics()


async def compute_analytics():
    timetable = await db.timetables.find_one({"active": True}, {"_id": 0})
    entries = timetable.get("entries", []) if timetable else []
    config = await read_config()
    total_slots = len(config.get("working_days", [])) * int(config.get("periods_per_day", 0))
    teachers = {t["id"]: t for t in await list_entities("teachers")}
    days = config.get("working_days", [])
    workload = []
    for teacher_id, count in Counter(e["teacher_id"] for e in entries).items():
        teacher = teachers.get(teacher_id, {})
        per_day = Counter(e["day"] for e in entries if e["teacher_id"] == teacher_id)
        workload.append({"id": teacher_id, "name": teacher.get("name") or next(e["teacher_name"] for e in entries if e["teacher_id"] == teacher_id), "lectures": count, "max_per_week": int(teacher.get("maximum_lectures_per_week") or 0), "max_per_day": int(teacher.get("maximum_lectures_per_day") or 0), "busiest_day": max(per_day, key=per_day.get) if per_day else "", "per_day": {day: per_day.get(day, 0) for day in days}, "labs": sum(1 for e in entries if e["teacher_id"] == teacher_id and e.get("requires_lab"))})
    workload.sort(key=lambda item: -item["lectures"])
    rooms = []
    for room_id, count in Counter(e["room_id"] for e in entries).items():
        sample = next(e for e in entries if e["room_id"] == room_id)
        rooms.append({"id": room_id, "name": sample["room_name"], "sessions": count, "total_slots": total_slots, "utilization": round(count / total_slots * 100) if total_slots else 0, "is_lab": bool(sample.get("requires_lab")), "divisions": sorted({e["division_name"] for e in entries if e["room_id"] == room_id})})
    rooms.sort(key=lambda item: -item["sessions"])
    return {"faculty_workload": workload, "room_utilization": rooms, "subject_distribution": [{"name": key, "sessions": value} for key, value in Counter(e["subject_name"] for e in entries).items()], "daily_density": [{"name": day, "sessions": sum(1 for e in entries if e["day"] == day), "labs": sum(1 for e in entries if e["day"] == day and e.get("requires_lab"))} for day in days], "quality_score": timetable.get("score", 0) if timetable else 0, "total_slots": total_slots, "total_sessions": len(entries), "timetable_name": timetable.get("name", "") if timetable else "", "hard_constraint_violations": timetable.get("hard_constraint_violations", 0) if timetable else 0, "college_name": config.get("college_name", ""), "academic_year": config.get("academic_year", ""), "timetable_id": timetable.get("id") if timetable else None}


@api.get("/analytics/report.pdf")
async def analytics_report(background: BackgroundTasks, _: dict = Depends(admin_user)):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    data = await compute_analytics()
    if not data["total_sessions"]:
        raise HTTPException(status_code=404, detail="Generate a timetable before exporting the analytics report")
    blue, teal, amber, ink, muted = colors.HexColor("#2563EB"), colors.HexColor("#14B8A6"), colors.HexColor("#F59E0B"), colors.HexColor("#0F172A"), colors.HexColor("#64748B")
    palette = [blue, teal, amber, colors.HexColor("#8B5CF6"), colors.HexColor("#EF4444"), colors.HexColor("#0EA5E9"), colors.HexColor("#84CC16"), colors.HexColor("#EC4899")]
    styles = getSampleStyleSheet()
    kicker = ParagraphStyle("kicker", parent=styles["Normal"], fontSize=7, textColor=blue, leading=9, spaceAfter=2)
    title = ParagraphStyle("title", parent=styles["Title"], fontSize=20, leading=24, alignment=0, textColor=ink, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=8.5, textColor=muted, leading=11)
    head = ParagraphStyle("head", parent=styles["Normal"], fontSize=9.5, leading=12, textColor=ink, fontName="Helvetica-Bold")
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7, leading=9, textColor=muted)
    workload, rooms, subjects, density = data["faculty_workload"][:8], data["room_utilization"][:8], data["subject_distribution"][:8], data["daily_density"]
    avg_util = round(sum(r["utilization"] for r in rooms) / len(rooms)) if rooms else 0
    overloaded = [w["name"] for w in data["faculty_workload"] if w["max_per_week"] and w["lectures"] / w["max_per_week"] > 0.85]
    busiest = max(density, key=lambda d: d["sessions"])["name"] if density else "—"
    quietest = min(density, key=lambda d: d["sessions"])["name"] if density else "—"

    def stat(label, value, color):
        d = Drawing(120, 46)
        d.add(Rect(0, 0, 120, 46, rx=8, ry=8, fillColor=colors.HexColor("#F8FAFC"), strokeColor=colors.HexColor("#E2E8F0"), strokeWidth=0.6))
        d.add(Rect(0, 40, 120, 6, rx=3, ry=3, fillColor=color, strokeColor=None))
        d.add(String(10, 26, label, fontName="Helvetica", fontSize=6.5, fillColor=muted))
        d.add(String(10, 9, str(value), fontName="Helvetica-Bold", fontSize=15, fillColor=ink))
        return d

    def hbar(rows, key, color_fn, max_value, width=234, height=190, suffix=""):
        d = Drawing(width, height)
        chart = HorizontalBarChart()
        chart.x, chart.y, chart.width, chart.height = 70, 8, width - 90, height - 16
        chart.data = [[r[key] for r in rows]]
        chart.categoryAxis.categoryNames = [r["name"][:18] for r in rows]
        chart.categoryAxis.labels.fontName, chart.categoryAxis.labels.fontSize, chart.categoryAxis.labels.fillColor = "Helvetica", 6.5, ink
        chart.categoryAxis.strokeColor = colors.HexColor("#E2E8F0")
        chart.valueAxis.valueMin, chart.valueAxis.valueMax = 0, max_value
        chart.valueAxis.labels.fontName, chart.valueAxis.labels.fontSize, chart.valueAxis.labels.fillColor = "Helvetica", 6, muted
        chart.valueAxis.strokeColor, chart.valueAxis.gridStrokeColor, chart.valueAxis.visibleGrid = colors.HexColor("#E2E8F0"), colors.HexColor("#F1F5F9"), 1
        chart.valueAxis.labelTextFormat = f"%d{suffix.replace('%', '%%')}"
        chart.bars.strokeColor = None
        chart.barWidth = 9
        chart.groupSpacing = 8
        for index, row in enumerate(rows):
            chart.bars[(0, index)].fillColor = color_fn(row)
        chart.barLabels.fontName, chart.barLabels.fontSize, chart.barLabels.fillColor, chart.barLabels.dx = "Helvetica-Bold", 6.5, ink, 12
        chart.barLabelFormat = f"%d{suffix.replace('%', '%%')}"
        d.add(chart)
        return d

    def pie(rows, width=234, height=190):
        d = Drawing(width, height)
        chart = Pie()
        chart.x, chart.y, chart.width, chart.height = 4, 30, 120, 120
        chart.data = [r["sessions"] for r in rows]
        chart.labels = None
        chart.slices.strokeColor, chart.slices.strokeWidth = colors.white, 1.2
        chart.sideLabels = 0
        for index in range(len(rows)):
            chart.slices[index].fillColor = palette[index % len(palette)]
        d.add(chart)
        total = sum(chart.data) or 1
        for index, row in enumerate(rows):
            y = height - 18 - index * 15
            d.add(Rect(134, y - 1, 7, 7, rx=2, ry=2, fillColor=palette[index % len(palette)], strokeColor=None))
            d.add(String(145, y, f"{row['name'][:20]} · {row['sessions']} ({round(row['sessions'] / total * 100)}%)", fontName="Helvetica", fontSize=6, fillColor=ink))
        return d

    def vbar(rows, width=234, height=190):
        d = Drawing(width, height)
        chart = VerticalBarChart()
        chart.x, chart.y, chart.width, chart.height = 28, 22, width - 40, height - 34
        chart.data = [[r["sessions"] for r in rows], [r["labs"] for r in rows]]
        chart.categoryAxis.categoryNames = [r["name"][:3] for r in rows]
        chart.categoryAxis.labels.fontName, chart.categoryAxis.labels.fontSize, chart.categoryAxis.labels.fillColor = "Helvetica", 6.5, ink
        chart.categoryAxis.strokeColor = colors.HexColor("#E2E8F0")
        chart.valueAxis.valueMin = 0
        chart.valueAxis.labels.fontName, chart.valueAxis.labels.fontSize, chart.valueAxis.labels.fillColor = "Helvetica", 6, muted
        chart.valueAxis.strokeColor, chart.valueAxis.gridStrokeColor, chart.valueAxis.visibleGrid = colors.HexColor("#E2E8F0"), colors.HexColor("#F1F5F9"), 1
        chart.bars[0].fillColor, chart.bars[1].fillColor = blue, teal
        chart.bars.strokeColor = None
        chart.groupSpacing, chart.barSpacing = 10, 2
        chart.barLabels.fontName, chart.barLabels.fontSize, chart.barLabels.fillColor, chart.barLabels.dy = "Helvetica-Bold", 6, ink, 4
        chart.barLabelFormat = "%d"
        d.add(chart)
        d.add(Rect(30, 4, 8, 6, fillColor=blue, strokeColor=None)); d.add(String(41, 4, "All sessions", fontName="Helvetica", fontSize=6, fillColor=muted))
        d.add(Rect(95, 4, 8, 6, fillColor=teal, strokeColor=None)); d.add(String(106, 4, "Laboratory sessions", fontName="Helvetica", fontSize=6, fillColor=muted))
        return d

    def card(kicker_text, heading, drawing, note):
        return Table([[Paragraph(kicker_text.upper(), kicker)], [Paragraph(heading, head)], [drawing], [Paragraph(note, small)]], colWidths=[250], style=TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#E2E8F0")), ("ROUNDEDCORNERS", [8, 8, 8, 8]), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (0, 0), 8), ("BOTTOMPADDING", (0, -1), (-1, -1), 8)]))

    max_lectures = max([w["lectures"] for w in workload] + [1])
    workload_chart = hbar(workload[::-1], "lectures", lambda r: amber if r["max_per_week"] and r["lectures"] / r["max_per_week"] > 0.85 else blue, max_lectures + 1)
    rooms_chart = hbar(rooms[::-1], "utilization", lambda r: teal if r["is_lab"] else blue, 100, suffix="%")
    generated = datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")
    stats_row = Table([[stat("Timetable quality", f"{data['quality_score']}/100", teal), stat("Hard violations", data["hard_constraint_violations"], blue if not data["hard_constraint_violations"] else amber), stat("Sessions / week", data["total_sessions"], blue), stat("Avg room utilization", f"{avg_util}%", amber)]], colWidths=[128] * 4, style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    grid = Table([
        [card("Faculty workload", "Lectures per teacher (per week)", workload_chart, f"Amber bars exceed 85% of the teacher's weekly limit. {('Watch: ' + ', '.join(overloaded)) if overloaded else 'No teacher is near their weekly limit.'}"), card("Room utilization", "% of weekly slots occupied", rooms_chart, f"Based on {data['total_slots']} teaching slots per week. Teal bars are laboratories.")],
        [card("Subject distribution", "Share of the timetable per subject", pie(subjects), f"{len(data['subject_distribution'])} subjects scheduled across {data['total_sessions']} sessions."), card("Daily density", "Sessions scheduled per working day", vbar(density), f"Busiest day is {busiest}; {quietest} is the lightest.")],
    ], colWidths=[258, 258], style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story = [Paragraph("SMART CLASSROOM · ANALYTICS REPORT", kicker), Paragraph(f"{data['college_name'] or 'College'} — timetable analytics", title), Paragraph(f"{data['timetable_name'] or 'Active timetable'} · Academic year {data['academic_year'] or '—'} · Generated {generated}", sub), Spacer(1, 10), stats_row, Spacer(1, 10), grid, Spacer(1, 4), Paragraph("Prepared automatically by Smart Classroom & Timetable Scheduler from the active published timetable. Figures reflect the schedule at the time of export.", small)]
    buffer = BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=10 * mm, title="Timetable analytics report").build(story)
    payload = buffer.getvalue()
    if data["timetable_id"]:
        background.add_task(archive_export, {"id": data["timetable_id"], "name": f"{data['timetable_name']} analytics report", "score": data["quality_score"]}, payload, "pdf", "application/pdf")
    return Response(content=payload, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="analytics-report-{datetime.now(timezone.utc).strftime("%Y%m%d")}.pdf"'})


DAY_ORDER = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


@api.post("/timetable/{timetable_id}/activate")
async def activate_timetable(timetable_id: str, background: BackgroundTasks, _: dict = Depends(admin_user)):
    result = await db.timetables.update_one({"id": timetable_id}, {"$set": {"active": True, "updated_at": now()}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Timetable not found")
    await db.timetables.update_many({"id": {"$ne": timetable_id}}, {"$set": {"active": False}})
    background.add_task(notify_teachers, timetable_id)
    return {"ok": True}


async def teacher_schedule(teacher_id):
    teacher = await db.teachers.find_one({"id": teacher_id}, {"_id": 0})
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found")
    timetable = await db.timetables.find_one({"active": True}, {"_id": 0}) or {}
    entries = [entry for entry in timetable.get("entries", []) if entry.get("teacher_id") == teacher_id]
    config = await read_config()
    return {"teacher": {"id": teacher["id"], "name": teacher["name"], "department": teacher.get("department", ""), "departments": teacher.get("departments", []), "employee_id": teacher.get("employee_id", ""), "photo_file_id": teacher.get("photo_file_id")}, "timetable_name": timetable.get("name", ""), "entries": entries, "working_days": config.get("working_days", []), "periods_per_day": int(config.get("periods_per_day", 6))}


@api.get("/public/teacher/{teacher_id}")
async def public_teacher_schedule(teacher_id: str):
    return await teacher_schedule(teacher_id)


@api.get("/me/schedule")
async def my_schedule(user: dict = Depends(teacher_user)):
    return await teacher_schedule(user["teacher_id"])


@api.get("/public/divisions")
async def public_divisions():
    divisions = await db.divisions.find({}, {"_id": 0, "id": 1, "name": 1, "department": 1, "semester": 1, "student_count": 1}).sort("name", 1).to_list(500)
    return divisions


@api.get("/public/division/{division_id}")
async def public_division_schedule(division_id: str):
    division = await db.divisions.find_one({"id": division_id}, {"_id": 0})
    if not division:
        raise HTTPException(status_code=404, detail="Division not found")
    timetable = await db.timetables.find_one({"active": True}, {"_id": 0}) or {}
    entries = [entry for entry in timetable.get("entries", []) if entry.get("division_id") == division_id]
    config = await read_config()
    return {"division": {"id": division["id"], "name": division["name"], "department": division.get("department", ""), "semester": division.get("semester"), "student_count": division.get("student_count")}, "timetable_name": timetable.get("name", ""), "updated_at": timetable.get("updated_at", ""), "entries": entries, "working_days": config.get("working_days", []), "periods_per_day": int(config.get("periods_per_day", 6)), "start_time": config.get("start_time", "09:00"), "period_duration": int(config.get("period_duration", 55))}


def timetable_grid(timetable):
    entries = timetable.get("entries", [])
    days = sorted({entry["day"] for entry in entries}, key=lambda day: DAY_ORDER.index(day) if day in DAY_ORDER else 99)
    periods = sorted({entry["period"] for entry in entries})
    divisions = sorted({entry["division_name"] for entry in entries})
    def cell(division, day, period):
        return next((entry for entry in entries if entry["division_name"] == division and entry["day"] == day and entry["period"] == period), None)
    return days, periods, divisions, cell


@api.get("/timetable/{timetable_id}/export/excel")
async def export_timetable_excel(timetable_id: str, background: BackgroundTasks, _: dict = Depends(admin_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    timetable = await db.timetables.find_one({"id": timetable_id}, {"_id": 0})
    if not timetable:
        raise HTTPException(status_code=404, detail="Timetable not found")
    days, periods, divisions, cell = timetable_grid(timetable)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for division in divisions:
        sheet = workbook.create_sheet((division or "Timetable")[:31])
        sheet.append([f"{timetable.get('name', 'Weekly Timetable')} — Division {division}"])
        sheet.append(["Period", *days])
        for period in periods:
            row = [f"P{period}"]
            for day in days:
                entry = cell(division, day, period)
                row.append(f"{entry.get('subject_code') or entry['subject_name']}\n{entry['teacher_name']}\n{entry['room_name']}" if entry else "Free")
            sheet.append(row)
        for item in sheet[2]:
            item.fill = header_fill
            item.font = header_font
        for row in sheet.iter_rows(min_row=3):
            for item in row:
                item.alignment = wrap
        sheet.column_dimensions["A"].width = 10
        for index in range(len(days)):
            sheet.column_dimensions[chr(66 + index)].width = 24
    buffer = BytesIO()
    workbook.save(buffer)
    background.add_task(archive_export, timetable, buffer.getvalue(), "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="timetable-{timetable_id}.xlsx"'})


@api.get("/timetable/{timetable_id}/export/pdf")
async def export_timetable_pdf(timetable_id: str, background: BackgroundTasks, _: dict = Depends(admin_user)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    timetable = await db.timetables.find_one({"id": timetable_id}, {"_id": 0})
    if not timetable:
        raise HTTPException(status_code=404, detail="Timetable not found")
    days, periods, divisions, cell = timetable_grid(timetable)
    styles = getSampleStyleSheet()
    cell_style = styles["Normal"]
    cell_style.fontSize = 7
    cell_style.leading = 9
    story = [Paragraph(timetable.get("name", "Weekly Timetable"), styles["Title"]), Paragraph(f"Quality score {timetable.get('score', 0)}/100 · Generated {str(timetable.get('created_at', ''))[:10]}", styles["Normal"]), Spacer(1, 14)]
    for division in divisions:
        story.append(Paragraph(f"Division {division}", styles["Heading2"]))
        data = [["Period", *days]]
        for period in periods:
            row = [f"P{period}"]
            for day in days:
                entry = cell(division, day, period)
                row.append(Paragraph(f"<b>{entry.get('subject_code') or entry['subject_name']}</b><br/>{entry['teacher_name']}<br/>{entry['room_name']}", cell_style) if entry else "—")
            data.append(row)
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")])]))
        story.extend([table, Spacer(1, 18)])
    buffer = BytesIO()
    SimpleDocTemplate(buffer, pagesize=landscape(A4), title=timetable.get("name", "Timetable")).build(story)
    background.add_task(archive_export, timetable, buffer.getvalue(), "pdf", "application/pdf")
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="timetable-{timetable_id}.pdf"'})


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=[os.environ["FRONTEND_URL"]], allow_methods=["*"], allow_headers=["*"])


@app.on_event("shutdown")
async def shutdown():
    client.close()